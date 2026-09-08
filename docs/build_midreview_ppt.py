"""Build the KNOPRIX Mid-Review presentation (Group A12, ADSA).

16:9 deck, saffron/charcoal theme matching the deployed app.
Run:  python build_midreview_ppt.py   (from this folder; needs python-pptx)
Output: Knoprix_MidReview_A12.pptx
"""
from __future__ import annotations

import os

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

# ── theme ──────────────────────────────────────────────────────────────────
BG        = RGBColor(0x17, 0x0D, 0x06)   # deep warm charcoal
PANEL     = RGBColor(0x22, 0x14, 0x0B)   # panel charcoal
PANEL2    = RGBColor(0x2B, 0x1A, 0x0E)
PRIMARY   = RGBColor(0xF9, 0x73, 0x16)   # saffron
ACCENT    = RGBColor(0xF5, 0x9E, 0x0B)   # amber
CREAM     = RGBColor(0xFF, 0xF7, 0xED)   # cream ivory
MUTED     = RGBColor(0xE8, 0xD5, 0xC0)   # soft tan text
FAINT     = RGBColor(0xB8, 0x9B, 0x7E)   # dimmed tan
GREEN     = RGBColor(0x4A, 0xD0, 0x6E)   # done
GREY      = RGBColor(0x8A, 0x74, 0x5E)   # not-yet

FONT = "Segoe UI"

OUT_DIR = os.path.dirname(os.path.abspath(__file__))
IMG_DIR = os.path.join(os.path.dirname(OUT_DIR), "..")   # project root parent (Downloads)
# screenshots live next to this workspace root
ROOT    = os.path.dirname(os.path.dirname(OUT_DIR))


def img(name: str) -> str:
    return os.path.join(ROOT, name)


# ── helpers ────────────────────────────────────────────────────────────────
def _set_font(run, size: float, bold=False, color=CREAM, italic=False, font=FONT):
    f = run.font
    f.size = Pt(size)
    f.bold = bold
    f.italic = italic
    f.color.rgb = color
    f.name = font


def textbox(slide, x, y, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
            line_spacing=1.0, space_after=6):
    """runs: list of paragraphs; each paragraph = list of (text, dict) tuples."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, para in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line_spacing
        p.space_after = Pt(space_after)
        for text, style in para:
            r = p.add_run()
            r.text = text
            _set_font(r, **style)
    return tb


def rect(slide, x, y, w, h, fill=PANEL, line=None, shape=MSO_SHAPE.RECTANGLE,
         line_w=0.75, radius=None):
    sp = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid()
        sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line
        sp.line.width = Pt(line_w)
    sp.shadow.inherit = False
    if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        try:
            sp.adjustments[0] = radius
        except Exception:
            pass
    return sp


def bg(slide):
    rect(slide, 0, 0, 13.333, 7.5, fill=BG)
    # subtle saffron glow top-left + bottom-right
    glow = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(-1.5), Inches(-1.8), Inches(4.2), Inches(4.2))
    glow.fill.solid(); glow.fill.fore_color.rgb = RGBColor(0x3A, 0x20, 0x0B)
    glow.line.fill.background(); glow.shadow.inherit = False
    glow2 = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(11.2), Inches(6.2), Inches(3.4), Inches(3.4))
    glow2.fill.solid(); glow2.fill.fore_color.rgb = RGBColor(0x33, 0x1E, 0x08)
    glow2.line.fill.background(); glow2.shadow.inherit = False


def header(slide, kicker, title, page):
    rect(slide, 0, 0, 0.16, 7.5, fill=PRIMARY)
    textbox(slide, 0.55, 0.32, 11.5, 0.3,
            [[(kicker.upper(), dict(size=11, bold=True, color=ACCENT))]])
    textbox(slide, 0.55, 0.58, 11.8, 0.75,
            [[(title, dict(size=27, bold=True, color=CREAM))]])
    rect(slide, 0.58, 1.32, 3.4, 0.035, fill=PRIMARY)
    textbox(slide, 12.4, 7.08, 0.7, 0.3,
            [[(str(page), dict(size=10, color=FAINT))]], align=PP_ALIGN.RIGHT)


def chip(slide, x, y, text, w=2.4, fill=PANEL2, color=CREAM, size=10.5, bold=True,
         line=RGBColor(0x4A, 0x30, 0x1A)):
    c = rect(slide, x, y, w, 0.34, fill=fill, line=line, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
    tf = c.text_frame
    tf.word_wrap = False
    tf.margin_left = tf.margin_right = Inches(0.08)
    tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = text
    _set_font(r, size, bold, color)
    return c


def arrow_right(slide, x, y, w=0.42, h=0.28, color=PRIMARY):
    a = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x), Inches(y), Inches(w), Inches(h))
    a.fill.solid(); a.fill.fore_color.rgb = color
    a.line.fill.background(); a.shadow.inherit = False
    return a


def arrow_down(slide, x, y, w=0.28, h=0.4, color=PRIMARY):
    a = slide.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, Inches(x), Inches(y), Inches(w), Inches(h))
    a.fill.solid(); a.fill.fore_color.rgb = color
    a.line.fill.background(); a.shadow.inherit = False
    return a


def card(slide, x, y, w, h, title, body, title_color=PRIMARY, size_t=13, size_b=10.5,
         fill=PANEL, line=None, body_color=MUTED):
    rect(slide, x, y, w, h, fill=fill, line=line, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08)
    textbox(slide, x + 0.18, y + 0.14, w - 0.36, h - 0.28,
            [[(title, dict(size=size_t, bold=True, color=title_color))],
             [(body, dict(size=size_b, color=body_color))]],
            line_spacing=1.05, space_after=4)


def set_cell(cell, text, size=9.5, bold=False, color=MUTED, fill=None, align=PP_ALIGN.LEFT):
    if fill is not None:
        cell.fill.solid(); cell.fill.fore_color.rgb = fill
    tf = cell.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.06)
    tf.margin_top = tf.margin_bottom = Inches(0.03)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run(); r.text = text
    _set_font(r, size, bold, color)


# ── build ──────────────────────────────────────────────────────────────────
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]

slide_no = 0

def new_slide(kicker, title):
    global slide_no
    slide_no += 1
    s = prs.slides.add_slide(BLANK)
    bg(s)
    header(s, kicker, title, slide_no)
    return s


# ══ 1 · TITLE ═════════════════════════════════════════════════════════════
s = prs.slides.add_slide(BLANK)
bg(s)
rect(s, 0, 0, 13.333, 0.14, fill=PRIMARY)
rect(s, 0, 7.36, 13.333, 0.14, fill=PRIMARY)

# decorative monogram
mono = rect(s, 0.85, 1.15, 1.05, 1.05, fill=PRIMARY, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.22)
mtf = mono.text_frame; mtf.vertical_anchor = MSO_ANCHOR.MIDDLE
mp = mtf.paragraphs[0]; mp.alignment = PP_ALIGN.CENTER
mr = mp.add_run(); mr.text = "K"; _set_font(mr, 40, True, RGBColor(0xFF, 0xFF, 0xFF))

textbox(s, 2.15, 1.18, 10.6, 1.0,
        [[("A DSA-POWERED KNOWLEDGE MANAGEMENT PLATFORM", dict(size=24, bold=True, color=ACCENT))],
         [("with Intelligent Document Interaction", dict(size=24, bold=True, color=CREAM))]])

textbox(s, 0.9, 2.65, 11.6, 0.5,
        [[("KNOPRIX  ·  Mid-Semester Review  ·  ~50% of final project", dict(size=15, bold=True, color=PRIMARY))]])

textbox(s, 0.9, 3.35, 11.6, 0.35,
        [[("Advanced Data Structures & Algorithms (ADSA)", dict(size=13, color=MUTED))]])
textbox(s, 0.9, 3.7, 11.6, 0.35,
        [[("Group A12", dict(size=13, bold=True, color=CREAM))]])

# member cards
members = [
    ("T.V.S.S Phanindra Guptha", "cb.ai.u4aid25055"),
    ("Athul V.R", "cb.ai.u4aid25005"),
    ("M. Vittal", "cb.ai.u4aid25033"),
    ("K. Nanda Kishore", "cb.ai.u4aid25027"),
]
x0, y0, cw, ch, gap = 0.9, 4.45, 2.78, 1.15, 0.22
for i, (name, roll) in enumerate(members):
    cx = x0 + i * (cw + gap)
    rect(s, cx, y0, cw, ch, fill=PANEL, line=RGBColor(0x4A, 0x30, 0x1A),
         shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.1)
    textbox(s, cx + 0.18, y0 + 0.18, cw - 0.36, 0.6,
            [[(name, dict(size=12.5, bold=True, color=CREAM))]])
    textbox(s, cx + 0.18, y0 + 0.58, cw - 0.36, 0.4,
            [[(roll, dict(size=11, color=ACCENT))]])

textbox(s, 0.9, 6.05, 11.6, 0.4,
        [[("Deployed live:  ", dict(size=12, color=FAINT)),
          ("knoprixv2midreview.netlify.app", dict(size=12, bold=True, color=CREAM)),
          ("   ·   API: knoprix-midreview-api.onrender.com", dict(size=12, bold=True, color=CREAM))]])

textbox(s, 0.9, 6.6, 11.6, 0.4,
        [[("Guided by: Faculty Mentor  ·  Dept. of AI & DS  ·  August 2026", dict(size=11, color=FAINT))]])

# ══ 2 · MOTIVATION & PROBLEM ═══════════════════════════════════════════════
s = new_slide("02 · Context", "Motivation & Problem Statement")

card(s, 0.55, 1.6, 6.0, 2.55, "Why this project", "Students, researchers and professionals "
     "work with vast amounts of information scattered across PDFs, presentations, "
     "documents, notes and images.", size_t=14, size_b=12)
textbox(s, 0.75, 2.42, 5.6, 1.6,
        [[("•  ", dict(size=11.5, color=ACCENT)), ("Content lives in many isolated formats", dict(size=11.5, color=MUTED))],
         [("•  ", dict(size=11.5, color=ACCENT)), ("Finding a passage later is slow & manual", dict(size=11.5, color=MUTED))],
         [("•  ", dict(size=11.5, color=ACCENT)), ("Key ideas get lost — nothing is linked", dict(size=11.5, color=MUTED))]],
        line_spacing=1.15, space_after=5)

card(s, 6.78, 1.6, 6.0, 2.55, "What exists today", "File managers, note apps and AI "
     "chatbots each solve one slice — but none unify organisation, retrieval and recall.",
     size_t=14, size_b=12)
textbox(s, 6.98, 2.42, 5.6, 1.6,
        [[("•  ", dict(size=11.5, color=ACCENT)), ("File explorers: no content-level search", dict(size=11.5, color=MUTED))],
         [("•  ", dict(size=11.5, color=ACCENT)), ("Note apps (Obsidian/Roam): manual linking", dict(size=11.5, color=MUTED))],
         [("•  ", dict(size=11.5, color=ACCENT)), ("AI agents: black-box, opaque, heavy compute", dict(size=11.5, color=MUTED))]],
        line_spacing=1.15, space_after=5)

pb = rect(s, 0.55, 4.45, 12.23, 2.35, fill=PANEL, line=RGBColor(0x4A, 0x30, 0x1A),
          shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.07)
textbox(s, 0.85, 4.7, 11.7, 0.35,
        [[("PROBLEM STATEMENT", dict(size=12, bold=True, color=PRIMARY))]])
textbox(s, 0.85, 5.12, 11.7, 1.5,
        [[("There is no unified, ", dict(size=14, color=MUTED)),
          ("transparent DSA-based platform", dict(size=14, bold=True, color=CREAM)),
          (" that organises multi-format documents, enables instant prefix "
           "autocomplete and ranked full-text retrieval, and provides indexed "
           "of important passages — all explainable, self-contained, and without "
           "black-box ML or external API keys.", dict(size=14, color=MUTED))]],
        line_spacing=1.2)

# ══ 3 · LITERATURE SURVEY ══════════════════════════════════════════════════
s = new_slide("03 · Literature Survey", "Literature Survey (tabular)")

rows = [
    ("Lewis, P. et al.", "Retrieval-Augmented Generation for Knowledge-Intensive NLP Tasks", "2020", "NeurIPS", "Parametric + non-parametric memory; retrieval grounds answers — but relies on heavy neural retrievers (black-box)"),
    ("Zobel, J. & Moffat, A.", "Inverted Files for Text Search Engines", "2006", "ACM Computing Surveys", "Canonical inverted-index design & compression — foundation of our full-text index"),
    ("Fredkin, E.", "Trie Memory", "1960", "Communications of the ACM", "Origin of the trie — O(L) prefix lookup; basis of our autocomplete"),
    ("Ilyas, I. F. et al.", "A Survey of Top-k Query Processing Techniques in Relational Databases", "2008", "ACM Computing Surveys", "Top-k with threshold/heap algorithms — motivates Min-Heap ranking (final review)"),
    ("Kahn, A. B.", "Topological Sorting of Large Networks", "1962", "Communications of the ACM", "Graph ordering — basis of our Knowledge Graph linking (final review)"),
    ("Brin, S. & Page, L.", "The Anatomy of a Large-Scale Hypertextual Web Search Engine", "1998", "Computer Networks & ISDN", "Web-scale retrieval architecture — context for ranked, snippet-based results"),
]
tbl_shape = s.shapes.add_table(len(rows) + 1, 5, Inches(0.55), Inches(1.6), Inches(12.23), Inches(4.0))
tbl = tbl_shape.table
tbl.columns[0].width = Inches(1.9)
tbl.columns[1].width = Inches(4.35)
tbl.columns[2].width = Inches(0.75)
tbl.columns[3].width = Inches(1.75)
tbl.columns[4].width = Inches(3.48)
hdr = ["Author(s)", "Title", "Year", "Source", "Relevance / Gap addressed"]
for j, htxt in enumerate(hdr):
    set_cell(tbl.cell(0, j), htxt, size=10.5, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF),
             fill=RGBColor(0x8A, 0x3D, 0x0D), align=PP_ALIGN.LEFT)
for i, row in enumerate(rows, start=1):
    fill = PANEL if i % 2 else PANEL2
    for j, val in enumerate(row):
        set_cell(tbl.cell(i, j), val, size=9, bold=(j == 0), color=CREAM if j in (0, 1) else MUTED, fill=fill)
    tbl.rows[i].height = Inches(0.62)
tbl.rows[0].height = Inches(0.34)

textbox(s, 0.55, 5.85, 12.2, 0.9,
        [[("Tools studied: ", dict(size=10.5, bold=True, color=ACCENT)),
          ("Obsidian (2020), Roam Research (2019), Google NotebookLM (2023) — powerful, but require "
           "manual linking or cloud AI; no transparent local DSA retrieval.", dict(size=10.5, color=FAINT))]],
        line_spacing=1.15)

# ══ 4 · RESEARCH GAP & NOVELTY ═════════════════════════════════════════════
s = new_slide("04 · Novelty", "Research Gap & Novelty of the Work")

textbox(s, 0.55, 1.55, 12.2, 0.35,
        [[("Research gaps identified", dict(size=13, bold=True, color=ACCENT))]])
gaps = [
    ("Black-box retrieval", "AI assistants fetch answers with opaque, unverifiable reasoning."),
    ("Manual knowledge linking", "Note tools require users to create every connection by hand."),
    ("No content-level search", "File managers index names, not the words inside documents."),
    ("External dependency", "Cloud/API-based tools need keys, accounts and internet."),
]
for i, (t, b) in enumerate(gaps):
    cx = 0.55 + i * 3.1
    card(s, cx, 1.95, 2.92, 1.7, t, b, size_t=12, size_b=9.5, fill=PANEL, line=RGBColor(0x4A, 0x30, 0x1A))

textbox(s, 0.55, 3.95, 12.2, 0.35,
        [[("Novelty of KNOPRIX (mid-review build)", dict(size=13, bold=True, color=ACCENT))]])
nov = [
    ("DSA-driven intelligence layer", "Trie + Inverted Index + bookmark collection replace black-box ML with transparent, explainable retrieval."),
    ("Indexed bookmark collection", "Hash Table + Doubly Linked List keeps bookmarks newest-first and supports average O(1) ID deletion."),
    ("Ranked search with scores", "Inverted Index returns results with match scores and snippets — instant, self-contained."),
    ("Trie autocomplete with frequency", "O(L) prefix suggestions ranked by real word frequency across your documents."),
    ("Multi-format workspace", "PDF, PPTX, DOCX, TXT/MD in isolated projects — one place to read everything."),
    ("Fully self-contained & secure", "JWT auth, owner-scoped routes (IDOR-proof), no API keys, deployed live on Netlify + Render."),
]
for i, (t, b) in enumerate(nov):
    col, row = i % 3, i // 3
    cx = 0.55 + col * 4.13
    cy = 4.35 + row * 1.42
    num = rect(s, cx, cy, 0.34, 0.34, fill=PRIMARY, shape=MSO_SHAPE.OVAL)
    ntf = num.text_frame; ntf.vertical_anchor = MSO_ANCHOR.MIDDLE
    np_ = ntf.paragraphs[0]; np_.alignment = PP_ALIGN.CENTER
    nr = np_.add_run(); nr.text = str(i + 1); _set_font(nr, 11, True, RGBColor(0xFF, 0xFF, 0xFF))
    textbox(s, cx + 0.46, cy - 0.02, 3.6, 1.3,
            [[(t, dict(size=11.5, bold=True, color=CREAM))],
             [(b, dict(size=9, color=MUTED))]], line_spacing=1.05, space_after=3)

# ══ 5 · OBJECTIVES ════════════════════════════════════════════════════════
s = new_slide("05 · Objectives", "Project Objectives (mid-review scope)")

objs = [
    ("Unified document management", "Project-based workspace for PDF, PPTX, DOCX and TXT/MD — organised folders with a Tree + DFS hierarchy view."),
    ("Trie-based instant autocomplete", "Live prefix suggestions in the navbar, ranked by word frequency across all documents (O(L) search)."),
    ("Inverted-Index ranked full-text search", "Type a phrase → scored results with snippets; click any result to jump straight into the document."),
    ("Hash Table + DLL bookmarking", "Every bookmark is indexed by ID and linked newest-first; add, lookup, and unlink are average O(1)."),
    ("Secure, deployable full-stack system", "JWT + bcrypt auth, rate-limited login, owner-scoped routes, SQLite persistence — deployed on Netlify + Render."),
]
y = 1.62
for i, (t, b) in enumerate(objs):
    num = rect(s, 0.55, y, 0.42, 0.42, fill=PRIMARY, shape=MSO_SHAPE.OVAL)
    ntf = num.text_frame; ntf.vertical_anchor = MSO_ANCHOR.MIDDLE
    np_ = ntf.paragraphs[0]; np_.alignment = PP_ALIGN.CENTER
    nr = np_.add_run(); nr.text = str(i + 1); _set_font(nr, 13, True, RGBColor(0xFF, 0xFF, 0xFF))
    textbox(s, 1.2, y - 0.04, 11.5, 0.85,
            [[(t, dict(size=13.5, bold=True, color=CREAM))],
             [(b, dict(size=11, color=MUTED))]], line_spacing=1.08, space_after=3)
    if i < len(objs) - 1:
        rect(s, 0.76, y + 0.5, 0.004, 0.5, fill=RGBColor(0x4A, 0x30, 0x1A))
    y += 1.12

textbox(s, 0.55, 7.0, 12.2, 0.4,
        [[("Final review extends scope: ", dict(size=10.5, bold=True, color=FAINT)),
          ("Knowledge Graph · Min-Heap top-k ranking · AI document agent · read-aloud", dict(size=10.5, color=FAINT))]])

# ══ 6 · WORKFLOW ══════════════════════════════════════════════════════════
s = new_slide("06 · Workflow", "Work-Flow Diagram")

def flow_box(x, y, w, h, title, sub=None, fill=PANEL, tcolor=CREAM):
    rect(s, x, y, w, h, fill=fill, line=RGBColor(0x5A, 0x3B, 0x20),
         shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.12)
    textbox(s, x + 0.1, y + 0.12, w - 0.2, 0.4,
            [[(title, dict(size=11.5, bold=True, color=tcolor))]], align=PP_ALIGN.CENTER)
    if sub:
        textbox(s, x + 0.1, y + 0.45, w - 0.2, 0.5,
                [[(sub, dict(size=8.5, color=FAINT))]], align=PP_ALIGN.CENTER)

# Row 1 — indexing pipeline
r1y = 1.75
bw, bh, gap = 2.32, 1.05, 0.42
b1x = 0.55
flow_box(b1x, r1y, bw, bh, "Upload documents", "PDF · PPTX · DOCX · TXT/MD")
flow_box(b1x + bw + gap, r1y, bw, bh, "Parse & extract", "pdf / pptx / docx parser")
flow_box(b1x + 2*(bw + gap), r1y, bw, bh, "Tokenize", "lowercase · stopwords · split")
flow_box(b1x + 3*(bw + gap), r1y, bw, bh, "Build DSA engine", "Trie + Inverted Index")
for i in range(3):
    arrow_right(s, b1x + (i + 1) * bw + i * gap + 0.02, r1y + 0.38)

# persist bar
py = 3.05
rect(s, b1x + 3*(bw+gap) - 0.2, py, 2.72, 0.62, fill=RGBColor(0x33, 0x20, 0x0E),
     line=RGBColor(0x5A, 0x3B, 0x20), shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.2)
textbox(s, b1x + 3*(bw+gap) - 0.05, py + 0.1, 2.42, 0.45,
        [[("Persist to SQLite  ·  dsa_indices (JSON)", dict(size=9.5, bold=True, color=ACCENT))]], align=PP_ALIGN.CENTER)
arrow_down(s, b1x + 3*(bw+gap) + 1.26 - 0.14, r1y + bh + 0.03)
arrow_down(s, b1x + 3*(bw+gap) + 1.26 - 0.14, py + 0.62)

# Row 2 — query pipeline
r2y = 4.15
flow_box(b1x, r2y, bw, bh, "Navbar query", "word → Trie · phrase → Inverted Index", fill=PANEL2)
flow_box(b1x + bw + gap, r2y, bw, bh, "Autocomplete / Search", "ranked results + snippets", fill=PANEL2)
flow_box(b1x + 2*(bw + gap), r2y, bw, bh, "Open document", "PDF / PPTX slide viewer", fill=PANEL2)
flow_box(b1x + 3*(bw + gap), r2y, bw, bh, "Bookmark Collection", "hash lookup · DLL order · O(1) avg", fill=PANEL2)
for i in range(3):
    arrow_right(s, b1x + (i + 1) * bw + i * gap + 0.02, r2y + 0.38)

arrow_down(s, b1x + 0.85, py + 0.62, h=0.46)
textbox(s, 0.55, 5.5, 12.2, 0.35,
        [[("Indices rebuild automatically on upload / delete  →  always fresh data", dict(size=10.5, color=FAINT))]])
textbox(s, 0.55, 5.85, 12.2, 0.35,
        [[("Every step is pure Python + SQLite — transparent, auditable, no external AI calls", dict(size=10.5, color=FAINT))]])

# ══ 7 · IMPLEMENTATION ════════════════════════════════════════════════════
s = new_slide("07 · Implementation", "Implementation Details")

textbox(s, 0.55, 1.55, 6.4, 0.3, [[("Tech stack", dict(size=12, bold=True, color=ACCENT))]])
stack = ["FastAPI · Python 3.13", "SQLite", "React 18 · Vite", "Tailwind CSS", "pdf.js · python-pptx · python-docx", "JWT + bcrypt · rate limiting", "Netlify + Render deploy"]
for i, t in enumerate(stack):
    col, row = i % 2, i // 2
    chip(s, 0.55 + col * 3.2, 1.9 + row * 0.44, t, w=3.05, size=10)

textbox(s, 7.4, 1.55, 5.4, 0.3, [[("DSA engine — where the code lives", dict(size=12, bold=True, color=ACCENT))]])
textbox(s, 7.4, 1.9, 5.4, 0.75,
        [[("backend/app/dsa/", dict(size=10.5, bold=True, color=CREAM)),
          (" — one file per data structure:", dict(size=10.5, color=MUTED))]], line_spacing=1.15)

dstructs = [
    ("trie.py", "Trie", "prefix tree with word frequencies", "insert/search/autocomplete O(L) · top keywords"),
    ("inverted_index.py", "Inverted Index", "word → {doc_id: count}", "build O(n) · lookup O(1) · ranked scores"),
    ("bookmark_collection.py", "Hash Table + DLL", "indexed newest-first nodes", "add/get/unlink O(1) avg · powers bookmarks"),
]
tbl_shape = s.shapes.add_table(4, 4, Inches(0.55), Inches(3.0), Inches(12.23), Inches(2.5))
tbl = tbl_shape.table
tbl.columns[0].width = Inches(2.35)
tbl.columns[1].width = Inches(2.6)
tbl.columns[2].width = Inches(3.6)
tbl.columns[3].width = Inches(3.68)
for j, htxt in enumerate(["File (app/dsa/)", "Structure", "Design", "Complexity / role"]):
    set_cell(tbl.cell(0, j), htxt, size=10, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF),
             fill=RGBColor(0x8A, 0x3D, 0x0D))
for i, (f, name, des, comp) in enumerate(dstructs, start=1):
    fill = PANEL if i % 2 else PANEL2
    set_cell(tbl.cell(i, 0), f, size=9, bold=True, color=ACCENT, fill=fill)
    set_cell(tbl.cell(i, 1), name, size=9, bold=True, color=CREAM, fill=fill)
    set_cell(tbl.cell(i, 2), des, size=9, color=MUTED, fill=fill)
    set_cell(tbl.cell(i, 3), comp, size=9, color=MUTED, fill=fill)
    tbl.rows[i].height = Inches(0.56)
tbl.rows[0].height = Inches(0.34)

textbox(s, 0.55, 5.85, 12.2, 1.0,
        [[("• ", dict(size=10.5, color=ACCENT)),
          ("Indices serialised to SQLite (", dict(size=10.5, color=MUTED)),
          ("dsa_indices", dict(size=10.5, bold=True, color=CREAM)),
          (") and rebuilt on upload/delete via ", dict(size=10.5, color=MUTED)),
          ("services/indexer.py", dict(size=10.5, bold=True, color=CREAM))],
         [("• ", dict(size=10.5, color=ACCENT)),
          ("Bookmarks stored in SQLite, ordered through the bookmark collection (", dict(size=10.5, color=MUTED)),
          ("routers/bookmarks.py", dict(size=10.5, bold=True, color=CREAM)),
          (") — create = add, delete = unlink", dict(size=10.5, color=MUTED))],
         [("• ", dict(size=10.5, color=ACCENT)),
          ("Security: JWT refresh flow, bcrypt hashing, rate-limited login, owner-scoped "
           "queries (IDOR-proof — attacker gets 404, verified by 16/16 API tests)", dict(size=10.5, color=MUTED))]],
        line_spacing=1.12, space_after=4)

# ══ 8 · OUTPUTS ═══════════════════════════════════════════════════════════
s = new_slide("08 · Outputs", "Outputs — live deployed app")

shots = [
    ("ppt-01-login.png", "Login — orange Saraswati theme"),
    ("ppt-03-dashboard.png", "Dashboard — MSA project, 8 documents"),
    ("ppt-04-trie-autocomplete.png", "Trie autocomplete — “signal” → freq-ranked"),
    ("ppt-05-inverted-index-search.png", "Inverted Index — ranked results + scores"),
]
for i, (fn, cap) in enumerate(shots):
    col, row = i % 2, i // 2
    cx = 0.55 + col * 6.3
    cy = 1.55 + row * 2.85
    rect(s, cx, cy, 6.1, 2.45, fill=PANEL, line=RGBColor(0x4A, 0x30, 0x1A),
         shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.05)
    try:
        s.shapes.add_picture(img(fn), Inches(cx + 0.08), Inches(cy + 0.08),
                             width=Inches(5.94))
    except Exception as e:
        textbox(s, cx + 0.2, cy + 0.9, 5.7, 0.6, [[(f"image missing: {fn}", dict(size=10, color=FAINT))]])
    textbox(s, cx + 0.1, cy + 2.12, 5.9, 0.3,
            [[(cap, dict(size=10.5, bold=True, color=CREAM))]], align=PP_ALIGN.CENTER)

textbox(s, 0.55, 7.02, 12.2, 0.35,
        [[("Demo account: demo@knoprix.io / Password123!  ·  Live at knoprixv2midreview.netlify.app", dict(size=10, color=FAINT))]])

# ══ 9 · OUTPUTS 2 (DSA modal + stack) ══════════════════════════════════════
s = new_slide("08 · Outputs", "Outputs — DSA visualizer & bookmark collection")

shots2 = [
    ("ppt-06-dsa-visualizer.png", "DSA Index Visualizer — Trie · Index · bookmark collection stats"),
    ("ppt-07-bookmarks-stack-full.png", "Bookmark collection drawer — 3 bookmarks, newest highlighted"),
]
for i, (fn, cap) in enumerate(shots2):
    cx = 0.55 + i * 6.3
    cy = 1.55
    rect(s, cx, cy, 6.1, 4.35, fill=PANEL, line=RGBColor(0x4A, 0x30, 0x1A),
         shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.05)
    try:
        s.shapes.add_picture(img(fn), Inches(cx + 0.08), Inches(cy + 0.08),
                             width=Inches(5.94))
    except Exception as e:
        textbox(s, cx + 0.2, cy + 1.8, 5.7, 0.6, [[(f"image missing: {fn}", dict(size=10, color=FAINT))]])
    textbox(s, cx + 0.1, cy + 3.98, 5.9, 0.3,
            [[(cap, dict(size=10.5, bold=True, color=CREAM))]], align=PP_ALIGN.CENTER)

textbox(s, 0.55, 6.15, 12.2, 0.9,
        [[("Live numbers at capture time: ", dict(size=11, bold=True, color=ACCENT)),
          ("2,978 Trie keywords · 2,978 index mappings · 3 bookmarks (newest = page 3)", dict(size=11, color=CREAM))],
         [("Rebuild one-click: re-tokenises every document and refreshes the index in place.", dict(size=10, color=FAINT))]],
        line_spacing=1.15, space_after=4)

# ══ 10 · TIMELINE ═════════════════════════════════════════════════════════
s = new_slide("09 · Work Progress", "Work Progress Timeline")

# progress bar 50%
textbox(s, 0.55, 1.55, 6.5, 0.3, [[("Mid-review milestone: ~50% complete", dict(size=13, bold=True, color=CREAM))]])
rect(s, 0.55, 1.95, 12.23, 0.5, fill=RGBColor(0x2B, 0x1A, 0x0E), shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
rect(s, 0.55, 1.95, 6.2, 0.5, fill=PRIMARY, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
textbox(s, 0.75, 2.0, 2.0, 0.4, [[("50% DONE", dict(size=11, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF)))]])

# done vs to-do
textbox(s, 0.55, 2.7, 6.1, 0.3, [[("✅ Completed (mid review)", dict(size=12, bold=True, color=GREEN))]])
done = [
    "Trie autocomplete + top keywords",
    "Inverted Index ranked full-text search",
    "Hash Table + Doubly Linked List bookmarks — visual drawer",
    "Auth (JWT + bcrypt), projects, uploads",
    "PDF / PPTX / DOCX / TXT readers",
    "Orange Saraswati theme, light/dark",
    "Public deploy — Netlify + Render",
    "16/16 API tests · frontend build clean",
]
for i, t in enumerate(done):
    col, row = i % 2, i // 2
    textbox(s, 0.55 + col * 3.05, 3.05 + row * 0.42, 3.0, 0.4,
            [[("✔ ", dict(size=10.5, bold=True, color=GREEN)), (t, dict(size=10.5, color=MUTED))]])

textbox(s, 7.05, 2.7, 5.7, 0.3, [[("⏳ Yet to do (final review)", dict(size=12, bold=True, color=GREY))]])
todo = [
    "Knowledge Graph (auto concept linking)",
    "Min-Heap top-k passage ranking",
    "AI document agent (chat over files)",
    "Read-aloud (speech synthesis)",
    "Concept linking visualisation",
]
for i, t in enumerate(todo):
    textbox(s, 7.05, 3.05 + i * 0.42, 5.7, 0.4,
            [[("◌ ", dict(size=10.5, bold=True, color=GREY)), (t, dict(size=10.5, color=FAINT))]])

# timeline bar
textbox(s, 0.55, 5.5, 12.2, 0.3, [[("Journey", dict(size=12, bold=True, color=ACCENT))]])
tl = [
    ("Zeroth review", "Jul 2026", "Problem selection, literature, proposal", "done", GREEN),
    ("Mid review", "Aug 2026", "3 structures live + deployed", "current", PRIMARY),
    ("Final review", "Nov 2026", "Knowledge Graph · Min-Heap · AI agent · read-aloud", "planned", GREY),
]
x = 0.55
for i, (name, when, what, state, col) in enumerate(tl):
    wdt = 3.9
    rect(s, x, 5.85, wdt, 1.05, fill=PANEL if state != "current" else RGBColor(0x3A, 0x20, 0x0B),
         line=col if state == "current" else RGBColor(0x4A, 0x30, 0x1A),
         shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08)
    dot = rect(s, x + 0.15, 6.0, 0.24, 0.24, fill=col, shape=MSO_SHAPE.OVAL)
    textbox(s, x + 0.5, 5.93, wdt - 0.6, 0.4,
            [[(name, dict(size=11.5, bold=True, color=CREAM)),
              ("   " + when, dict(size=9, color=FAINT))]])
    textbox(s, x + 0.5, 6.3, wdt - 0.6, 0.55,
            [[(what, dict(size=8.5, color=MUTED))]], line_spacing=1.05)
    x += wdt + 0.27

# ══ 11 · REFERENCES ═══════════════════════════════════════════════════════
s = new_slide("10 · References", "References")

refs = [
    ("Lewis, P., Perez, E., Piktus, A., et al.", "“Retrieval-Augmented Generation for Knowledge-Intensive NLP Tasks.” Advances in Neural Information Processing Systems (NeurIPS), 2020."),
    ("Zobel, J. & Moffat, A.", "“Inverted Files for Text Search Engines.” ACM Computing Surveys, 38(2), 2006."),
    ("Fredkin, E.", "“Trie Memory.” Communications of the ACM, 3(9), 1960."),
    ("Ilyas, I. F., Beskales, G., Soliman, M. A.", "“A Survey of Top-k Query Processing Techniques in Relational Databases.” ACM Computing Surveys, 40(4), 2008."),
    ("Kahn, A. B.", "“Topological Sorting of Large Networks.” Communications of the ACM, 5(11), 1962."),
    ("Brin, S. & Page, L.", "“The Anatomy of a Large-Scale Hypertextual Web Search Engine.” Computer Networks and ISDN Systems, 30(1–7), 1998."),
    ("Cormen, T. H., Leiserson, C. E., Rivest, R. L., Stein, C.", "Introduction to Algorithms, 3rd ed., MIT Press, 2009 (Trie, heaps, hash tables)."),
    ("Tools reviewed", "Obsidian (Dynalist, 2020) · Roam Research (2019) · Google NotebookLM (2023)."),
]
y = 1.6
for i, (a, t) in enumerate(refs):
    textbox(s, 0.55, y, 0.5, 0.4, [[(f"{i+1:02d}", dict(size=12, bold=True, color=ACCENT))]])
    textbox(s, 1.05, y - 0.02, 11.7, 0.6,
            [[(a + " — ", dict(size=11, bold=True, color=CREAM)),
              (t, dict(size=11, color=MUTED))]], line_spacing=1.05)
    y += 0.63

# ══ THANK YOU ═════════════════════════════════════════════════════════════
s = prs.slides.add_slide(BLANK)
bg(s)
rect(s, 0, 0, 13.333, 0.14, fill=PRIMARY)
rect(s, 0, 7.36, 13.333, 0.14, fill=PRIMARY)
textbox(s, 1.2, 2.6, 11.0, 0.9,
        [[("Thank You", dict(size=44, bold=True, color=CREAM))]], align=PP_ALIGN.CENTER)
textbox(s, 1.2, 3.6, 11.0, 0.5,
        [[("Questions & feedback are welcome", dict(size=15, color=MUTED))]], align=PP_ALIGN.CENTER)
for i, (name, roll) in enumerate(members):
    cx = 1.55 + i * 2.72
    textbox(s, cx, 4.5, 2.5, 0.4, [[(name, dict(size=11, bold=True, color=CREAM))]], align=PP_ALIGN.CENTER)
    textbox(s, cx, 4.88, 2.5, 0.35, [[(roll, dict(size=10, color=ACCENT))]], align=PP_ALIGN.CENTER)
textbox(s, 1.2, 5.7, 11.0, 0.4,
        [[("KNOPRIX · A DSA-Powered Knowledge Management Platform · Group A12 · ADSA", dict(size=11, color=FAINT))]],
        align=PP_ALIGN.CENTER)

# ── save ───────────────────────────────────────────────────────────────────
out = os.path.join(OUT_DIR, "Knoprix_MidReview_A12.pptx")
prs.save(out)
print(f"Saved: {out}")
print(f"Slides: {len(prs.slides)}")
