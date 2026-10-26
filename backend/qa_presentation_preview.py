"""Verify accurate PPTX PDF previews using the deployment LibreOffice binary."""
import shutil
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from pptx import Presentation
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.util import Inches
from pypdf import PdfReader

from app.services import presentation_preview


BACKEND_DIR = Path(__file__).resolve().parent
LIBREOFFICE_AVAILABLE = bool(
    shutil.which("libreoffice") or shutil.which("soffice")
)


@unittest.skipUnless(LIBREOFFICE_AVAILABLE, "LibreOffice is required for PPTX preview QA")
class PresentationPreviewTests(unittest.TestCase):
    def test_flowchart_shapes_and_text_render_to_a_cached_pdf_page(self):
        with tempfile.TemporaryDirectory(dir=BACKEND_DIR) as temporary_directory:
            work_dir = Path(temporary_directory)
            source = work_dir / "flowchart.pptx"
            presentation = Presentation()
            slide = presentation.slides.add_slide(presentation.slide_layouts[6])
            start = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE,
                Inches(1),
                Inches(1),
                Inches(3),
                Inches(0.8),
            )
            start.text = "Start process"
            decision = slide.shapes.add_shape(
                MSO_SHAPE.DIAMOND,
                Inches(1.6),
                Inches(2.3),
                Inches(2),
                Inches(1),
            )
            decision.text = "Decision"
            slide.shapes.add_connector(
                MSO_CONNECTOR.STRAIGHT,
                Inches(2.5),
                Inches(1.8),
                Inches(2.5),
                Inches(2.3),
            )
            presentation.save(source)

            cache_dir = work_dir / "previews"
            with patch.object(presentation_preview, "PREVIEW_CACHE_DIR", cache_dir):
                preview = presentation_preview.render_presentation_preview(
                    source, "document-1"
                )
                pages = PdfReader(preview).pages
                text = pages[0].extract_text()
                cached_again = presentation_preview.render_presentation_preview(
                    source, "document-1"
                )

            self.assertEqual(len(pages), 1)
            normalized_text = "".join(text.split())
            self.assertIn("Startprocess", normalized_text)
            self.assertIn("Decision", normalized_text)
            self.assertEqual(cached_again, preview)
            self.assertTrue(preview.is_file())


if __name__ == "__main__":
    unittest.main(verbosity=2)
