"""Local regression checks for content-based, bounded document uploads."""
import asyncio
import io
import sqlite3
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch

from docx import Document
from fastapi import HTTPException
from fastapi.datastructures import UploadFile
from fastapi.responses import FileResponse
from pptx import Presentation
from pypdf import PdfWriter
from starlette.datastructures import Headers

from app import database
from app.routers import documents


USER = {"id": "user-1"}
PROJECT_ID = "project-1"


def _pdf_bytes():
    writer = PdfWriter()
    writer.add_blank_page(width=72, height=72)
    output = io.BytesIO()
    writer.write(output)
    return output.getvalue()


def _docx_bytes():
    document = Document()
    document.add_paragraph("Upload validation")
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def _pptx_bytes():
    presentation = Presentation()
    presentation.slides.add_slide(presentation.slide_layouts[6])
    output = io.BytesIO()
    presentation.save(output)
    return output.getvalue()


def _image_bytes(image_format):
    from PIL import Image

    output = io.BytesIO()
    Image.new("RGB", (2, 2), "gold").save(output, format=image_format)
    return output.getvalue()


def _zip_bytes(entries):
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, content in entries:
            archive.writestr(name, content)
    return output.getvalue()


class UploadValidationTests(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp_dir.cleanup)
        self.upload_dir = Path(self.temp_dir.name) / "uploads"
        self.db = sqlite3.connect(":memory:")
        self.db.row_factory = sqlite3.Row
        self.db.execute("PRAGMA foreign_keys = ON")
        self.db.executescript(database.SCHEMA)
        self.db.execute(
            "INSERT INTO users (id, email, password_hash, full_name) "
            "VALUES (?, ?, ?, ?)",
            ("user-1", "owner@example.com", "hash", "Owner"),
        )
        self.db.execute(
            "INSERT INTO projects (id, user_id, name) VALUES (?, ?, ?)",
            (PROJECT_ID, USER["id"], "Research"),
        )
        self.db.commit()
        self.addCleanup(self.db.close)
        self.upload_dir_patch = patch.object(documents, "UPLOAD_DIR", self.upload_dir)
        self.storage_patch = patch.object(documents.storage, "enabled", return_value=False)
        self.index_patch = patch.object(documents, "rebuild_project_indices")
        self.upload_dir_patch.start()
        self.storage_patch.start()
        self.index_patch.start()
        self.addCleanup(self.upload_dir_patch.stop)
        self.addCleanup(self.storage_patch.stop)
        self.addCleanup(self.index_patch.stop)

    def upload(self, filename, content, content_type="application/octet-stream"):
        upload = UploadFile(
            file=io.BytesIO(content),
            filename=filename,
            headers=Headers({"content-type": content_type}),
        )
        return asyncio.run(
            documents.upload_document(PROJECT_ID, upload, USER, self.db)
        )

    def test_accepts_real_supported_files_and_rejects_legacy_powerpoint(self):
        valid_files = (
            ("paper.pdf", _pdf_bytes()),
            ("slides.pptx", _pptx_bytes()),
            ("notes.docx", _docx_bytes()),
            ("notes.txt", "plain text".encode()),
            ("notes.md", "# Notes\n".encode()),
            ("scan.png", _image_bytes("PNG")),
            ("photo.jpg", _image_bytes("JPEG")),
            ("photo.jpeg", _image_bytes("JPEG")),
            ("diagram.svg", b'<svg xmlns="http://www.w3.org/2000/svg"/>'),
        )

        for filename, content in valid_files:
            with self.subTest(filename=filename):
                result = self.upload(filename, content)
                self.assertEqual(result["document"]["file_type"], Path(filename).suffix[1:])
        with self.assertRaises(HTTPException) as raised:
            self.upload("old.ppt", b"not a supported format")
        self.assertEqual(raised.exception.status_code, 415)

    def test_rejects_mismatched_content_and_malformed_files(self):
        for filename, content in (
            ("spoofed.pdf", b"ordinary text"),
            ("truncated.pdf", b"%PDF-1.7"),
            ("truncated.png", b"\x89PNG\r\n\x1a\n"),
            ("truncated.jpg", b"\xff\xd8\xff"),
            ("broken.docx", b"not a zip archive"),
        ):
            with self.subTest(filename=filename):
                with self.assertRaises(HTTPException):
                    self.upload(filename, content)

    def test_ignores_client_mime_and_sanitizes_display_filename(self):
        result = self.upload(
            "..\\..\\notes\r\n.txt",
            "valid text".encode(),
            content_type="application/pdf",
        )
        document = result["document"]
        self.assertEqual(document["file_name"], "notes.txt")
        self.assertEqual(Path(document["file_path"]).parent, self.upload_dir)

    def test_empty_and_oversized_uploads_are_removed(self):
        with self.assertRaises(HTTPException) as empty:
            self.upload("empty.txt", b"")
        self.assertEqual(empty.exception.status_code, 422)

        with patch.object(documents.cfg, "MAX_UPLOAD_BYTES", 8):
            with self.assertRaises(HTTPException) as oversized:
                self.upload("large.txt", b"x" * 9)
        self.assertEqual(oversized.exception.status_code, 413)
        self.assertEqual(list(self.upload_dir.glob("*")), [])

    def test_bounds_office_archive_expansion_and_rejects_unsafe_svg(self):
        office = _zip_bytes(
            (
                ("[Content_Types].xml", b"<Types/>"),
                ("word/document.xml", b"x" * 128),
            )
        )
        path = self.upload_dir / "compressed.docx"
        self.upload_dir.mkdir(parents=True)
        path.write_bytes(office)
        with patch.object(documents, "MAX_ARCHIVE_UNCOMPRESSED_BYTES", 64):
            with self.assertRaises(HTTPException) as oversized:
                documents._validate_upload_file(path, ".docx")
        self.assertEqual(oversized.exception.status_code, 413)

        active_svg = b'<svg xmlns="http://www.w3.org/2000/svg"><script>alert(1)</script></svg>'
        with self.assertRaises(HTTPException) as unsafe:
            self.upload("active.svg", active_svg)
        self.assertEqual(unsafe.exception.status_code, 422)

    def test_uploaded_file_download_is_owner_scoped(self):
        result = self.upload("notes.txt", b"private note")
        document_id = result["document"]["id"]
        response = documents.stream_document(document_id, USER, self.db)
        self.assertIsInstance(response, FileResponse)
        self.assertIn("inline", response.headers["content-disposition"])
        self.assertEqual(response.headers["x-content-type-options"], "nosniff")

        with self.assertRaises(HTTPException) as unauthorized:
            documents.stream_document(document_id, {"id": "other-user"}, self.db)
        self.assertEqual(unauthorized.exception.status_code, 404)

    def test_svg_is_download_only_and_storage_keys_ignore_upload_names(self):
        svg = b'<svg xmlns="http://www.w3.org/2000/svg"><rect width="1" height="1"/></svg>'
        result = self.upload("diagram.svg", svg)
        local_response = documents.stream_document(
            result["document"]["id"], USER, self.db
        )
        self.assertIn("attachment", local_response.headers["content-disposition"])
        self.assertEqual(local_response.headers["x-content-type-options"], "nosniff")

        with (
            patch.object(documents.storage, "enabled", return_value=True),
            patch.object(documents.storage, "upload") as upload,
            patch.object(documents.storage, "stream", return_value=iter((svg,))),
        ):
            stored = self.upload(r"..\..\diagram.svg", svg)
            object_response = documents.stream_document(
                stored["document"]["id"], USER, self.db
            )
        key = upload.call_args.args[1]
        content_type = upload.call_args.args[2]
        self.assertEqual(stored["document"]["file_name"], "diagram.svg")
        self.assertRegex(key, r"^documents/[0-9a-f-]+/[0-9a-f]{32}\.svg$")
        self.assertEqual(content_type, "application/octet-stream")
        self.assertIn("attachment", object_response.headers["content-disposition"])
        self.assertEqual(object_response.headers["content-type"], "application/octet-stream")
        self.assertEqual(object_response.headers["x-content-type-options"], "nosniff")


if __name__ == "__main__":
    unittest.main(verbosity=2)
